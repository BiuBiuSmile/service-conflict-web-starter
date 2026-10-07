import json
import hashlib
import hmac
import secrets
import time
import os
import re
import shutil
import subprocess
import threading
import webbrowser
from pathlib import Path
from functools import wraps
from datetime import date, datetime, timedelta
import tkinter as tk
from tkinter import messagebox

import requests
import websocket
from openpyxl import load_workbook
from flask import Flask, jsonify, request
from flask_cors import CORS

VERSION = "0.5.6"
PORT = 8765
WEBSITE_URL = "https://biubiusmile.github.io/service-conflict-web-starter/"

app = Flask(__name__)

ALLOWED_ORIGINS = [
    "https://biubiusmile.github.io",
    "https://chkia.dev",
    "http://localhost",
    "http://127.0.0.1",
]

CORS(
    app,
    resources={r"/*": {"origins": ALLOWED_ORIGINS}},
    supports_credentials=False,
)

SYSTEMS = {
    "compal": {
        "name": "仁寶",
        "url": "https://luna.compal-health.com/login",
        "host": "luna.compal-health.com",
        "port": 9360,
    },
    "lcms": {
        "name": "照管",
        "url": "https://csms.mohw.gov.tw/lcms/",
        "host": "csms.mohw.gov.tw",
        "port": 9460,
    },
}


AUTH_ITERATIONS = 210000
AUTH_TOKEN_TTL = 12 * 60 * 60
AUTH_LOCK = threading.Lock()
AUTH_SESSIONS = {}


def auth_file():
    return base_dir() / "users.json"


def load_users():
    path = auth_file()
    if not path.exists():
        return {}
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
        return data if isinstance(data, dict) else {}
    except Exception:
        return {}


def save_users(users):
    path = auth_file()
    tmp = path.with_suffix(".tmp")
    tmp.write_text(
        json.dumps(users, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    tmp.replace(path)


def hash_password(password, salt_hex=None):
    if salt_hex:
        salt = bytes.fromhex(salt_hex)
    else:
        salt = secrets.token_bytes(16)

    digest = hashlib.pbkdf2_hmac(
        "sha256",
        password.encode("utf-8"),
        salt,
        AUTH_ITERATIONS,
    )
    return salt.hex(), digest.hex()


def verify_password(password, salt_hex, hash_hex):
    try:
        _, candidate = hash_password(password, salt_hex)
        return hmac.compare_digest(candidate, hash_hex)
    except Exception:
        return False


def issue_token(username, role):
    token = secrets.token_urlsafe(32)
    AUTH_SESSIONS[token] = {
        "username": username,
        "role": role,
        "expires": time.time() + AUTH_TOKEN_TTL,
    }
    return token


def get_auth_context():
    header = request.headers.get("Authorization", "")
    if not header.startswith("Bearer "):
        return None

    token = header[7:].strip()
    session = AUTH_SESSIONS.get(token)
    if not session:
        return None

    if session.get("expires", 0) < time.time():
        AUTH_SESSIONS.pop(token, None)
        return None

    users = load_users()
    user = users.get(session.get("username"))
    if not user or not user.get("enabled", True):
        AUTH_SESSIONS.pop(token, None)
        return None

    return {
        "token": token,
        "username": session["username"],
        "role": user.get("role", "user"),
    }


def require_auth(fn):
    @wraps(fn)
    def wrapper(*args, **kwargs):
        ctx = get_auth_context()
        if not ctx:
            return jsonify(error="尚未登入或登入已逾時"), 401
        request.auth = ctx
        return fn(*args, **kwargs)
    return wrapper


def require_admin(fn):
    @wraps(fn)
    def wrapper(*args, **kwargs):
        ctx = get_auth_context()
        if not ctx:
            return jsonify(error="尚未登入或登入已逾時"), 401
        if ctx.get("role") != "admin":
            return jsonify(error="需要管理者權限"), 403
        request.auth = ctx
        return fn(*args, **kwargs)
    return wrapper


@app.get("/auth/status")
def auth_status():
    try:
        users = load_users()
        return jsonify(
            configured=bool(users),
            logged_in=bool(get_auth_context()),
            version=VERSION,
        )
    except Exception as e:
        return jsonify(
            configured=False,
            logged_in=False,
            version=VERSION,
            error=f"{type(e).__name__}: {e}",
        ), 200


@app.post("/auth/setup")
def auth_setup():
    payload = request.get_json(silent=True) or {}
    username = str(payload.get("username") or "").strip()
    password = str(payload.get("password") or "")

    if len(username) < 3:
        return jsonify(error="管理者帳號至少 3 個字元"), 400
    if len(password) < 8:
        return jsonify(error="密碼至少需要 8 個字元"), 400

    with AUTH_LOCK:
        users = load_users()

        if users:
            # 如果首次設定其實已在前一瞬間成功（例如重複點擊、
            # 頁面重送），且輸入的是同一組管理者帳密，
            # 直接視為登入成功，不再回 409。
            existing = users.get(username)
            if (
                existing
                and existing.get("role") == "admin"
                and existing.get("enabled", True)
                and verify_password(
                    password,
                    existing.get("salt", ""),
                    existing.get("password_hash", ""),
                )
            ):
                token = issue_token(username, "admin")
                return jsonify(
                    ok=True,
                    token=token,
                    user={"username": username, "role": "admin"},
                    already_configured=True,
                )

            return jsonify(
                error="管理者帳號已建立，請改用登入。",
                code="ALREADY_CONFIGURED",
            ), 409

        salt, password_hash = hash_password(password)
        users[username] = {
            "role": "admin",
            "enabled": True,
            "salt": salt,
            "password_hash": password_hash,
            "created_at": datetime.now().isoformat(timespec="seconds"),
        }
        save_users(users)

    token = issue_token(username, "admin")
    return jsonify(
        ok=True,
        token=token,
        user={"username": username, "role": "admin"},
    )


@app.post("/auth/login")
def auth_login():
    started = time.time()

    try:
        payload = request.get_json(silent=True) or {}
        username = str(payload.get("username") or "").strip()
        password = str(payload.get("password") or "")

        if not username or not password:
            return jsonify(error="請輸入帳號及密碼"), 400

        users = load_users()
        user = users.get(username)

        if not user:
            return jsonify(error="帳號或密碼錯誤"), 401

        if not user.get("enabled", True):
            return jsonify(error="此帳號已停用，請聯絡管理者"), 403

        if not verify_password(
            password,
            user.get("salt", ""),
            user.get("password_hash", ""),
        ):
            return jsonify(error="帳號或密碼錯誤"), 401

        role = user.get("role", "user")
        token = issue_token(username, role)

        return jsonify(
            ok=True,
            token=token,
            user={"username": username, "role": role},
            elapsed_ms=int((time.time() - started) * 1000),
        )

    except Exception as e:
        return jsonify(
            error=f"登入處理失敗：{type(e).__name__}: {e}",
        ), 500


@app.post("/auth/logout")
@require_auth
def auth_logout():
    token = request.auth.get("token")
    AUTH_SESSIONS.pop(token, None)
    return jsonify(ok=True)


@app.get("/auth/me")
@require_auth
def auth_me():
    return jsonify(
        username=request.auth["username"],
        role=request.auth["role"],
    )


@app.get("/admin/users")
@require_admin
def admin_users():
    users = load_users()
    rows = []
    for username, user in sorted(users.items()):
        rows.append({
            "username": username,
            "role": user.get("role", "user"),
            "enabled": bool(user.get("enabled", True)),
            "created_at": user.get("created_at", ""),
        })
    return jsonify(users=rows)


@app.post("/admin/users")
@require_admin
def admin_create_user():
    payload = request.get_json(silent=True) or {}
    username = str(payload.get("username") or "").strip()
    password = str(payload.get("password") or "")
    role = str(payload.get("role") or "user").strip()

    if len(username) < 3:
        return jsonify(error="帳號至少 3 個字元"), 400
    if len(password) < 8:
        return jsonify(error="密碼至少需要 8 個字元"), 400
    if role not in ("user", "admin"):
        return jsonify(error="無效的角色"), 400

    with AUTH_LOCK:
        users = load_users()
        if username in users:
            return jsonify(error="帳號已存在"), 409

        salt, password_hash = hash_password(password)
        users[username] = {
            "role": role,
            "enabled": True,
            "salt": salt,
            "password_hash": password_hash,
            "created_at": datetime.now().isoformat(timespec="seconds"),
        }
        save_users(users)

    return jsonify(ok=True)


@app.patch("/admin/users/<username>")
@require_admin
def admin_update_user(username):
    payload = request.get_json(silent=True) or {}

    with AUTH_LOCK:
        users = load_users()
        user = users.get(username)
        if not user:
            return jsonify(error="找不到帳號"), 404

        if "enabled" in payload:
            enabled = bool(payload.get("enabled"))
            if username == request.auth["username"] and not enabled:
                return jsonify(error="不能停用目前登入中的管理者帳號"), 400
            user["enabled"] = enabled

        if "role" in payload:
            role = str(payload.get("role") or "")
            if role not in ("user", "admin"):
                return jsonify(error="無效的角色"), 400
            if username == request.auth["username"] and role != "admin":
                return jsonify(error="不能移除自己的管理者權限"), 400
            user["role"] = role

        password = str(payload.get("password") or "")
        if password:
            if len(password) < 8:
                return jsonify(error="密碼至少需要 8 個字元"), 400
            salt, password_hash = hash_password(password)
            user["salt"] = salt
            user["password_hash"] = password_hash

        users[username] = user
        save_users(users)

    return jsonify(ok=True)


@app.delete("/admin/users/<username>")
@require_admin
def admin_delete_user(username):
    if username == request.auth["username"]:
        return jsonify(error="不能刪除目前登入中的管理者帳號"), 400

    with AUTH_LOCK:
        users = load_users()
        if username not in users:
            return jsonify(error="找不到帳號"), 404
        users.pop(username, None)
        save_users(users)

    return jsonify(ok=True)



@app.after_request
def add_private_network_headers(response):
    if request.headers.get("Access-Control-Request-Private-Network") == "true":
        response.headers["Access-Control-Allow-Private-Network"] = "true"
    response.headers["Cache-Control"] = "no-store"
    return response

def base_dir():
    base = Path(os.environ.get("LOCALAPPDATA") or Path.home() / "AppData" / "Local")
    p = base / "ServiceConflictConnector"
    p.mkdir(parents=True, exist_ok=True)
    return p

def profile_dir(system):
    p = base_dir() / "ChromeProfiles" / system
    p.mkdir(parents=True, exist_ok=True)
    return p

def find_chrome():
    candidates = [
        Path(os.environ.get("PROGRAMFILES", "")) / "Google/Chrome/Application/chrome.exe",
        Path(os.environ.get("PROGRAMFILES(X86)", "")) / "Google/Chrome/Application/chrome.exe",
        Path(os.environ.get("LOCALAPPDATA", "")) / "Google/Chrome/Application/chrome.exe",
    ]
    for c in candidates:
        if c.is_file():
            return c
    found = shutil.which("chrome") or shutil.which("chrome.exe")
    return Path(found) if found else None

def already_running():
    try:
        r = requests.get(f"http://127.0.0.1:{PORT}/health", timeout=0.7)
        return r.ok
    except Exception:
        return False

def cdp_json(system, path="/json/list"):
    cfg = SYSTEMS[system]
    r = requests.get(f"http://127.0.0.1:{cfg['port']}{path}", timeout=2)
    r.raise_for_status()
    return r.json()

def cdp_available(system):
    try:
        v = cdp_json(system, "/json/version")
        return bool(v.get("webSocketDebuggerUrl"))
    except Exception:
        return False

def launch_login(system):
    cfg = SYSTEMS[system]
    chrome = find_chrome()
    if chrome is None:
        raise RuntimeError("找不到 Google Chrome，請先安裝 Chrome。")

    if cdp_available(system):
        return False

    cmd = [
        str(chrome),
        f"--remote-debugging-port={cfg['port']}",
        f"--user-data-dir={profile_dir(system)}",
        "--profile-directory=Default",
        "--disable-notifications",
        "--remote-allow-origins=*",
        "--no-first-run",
        "--no-default-browser-check",
        "--start-maximized",
        cfg["url"],
    ]
    subprocess.Popen(
        cmd,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
        creationflags=(subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0),
    )
    return True

def browser_status(system):
    cfg = SYSTEMS[system]

    try:
        if not cdp_available(system):
            return {
                "running": False,
                "logged_in": False,
                "message": "專用 Chrome 尚未啟動。",
                "url": "",
            }

        targets = cdp_json(system, "/json/list")
        pages = [
            x for x in targets
            if isinstance(x, dict) and x.get("type") == "page"
        ]
        related = [
            x for x in pages
            if cfg["host"] in str(x.get("url") or "")
        ]

        if not related:
            return {
                "running": True,
                "logged_in": False,
                "message": f"{cfg['name']} 官方頁面已關閉。",
                "url": "",
            }

        url = str(related[0].get("url") or "")

        if system == "lcms":
            try:
                ok, message = probe_lcms()
                return {
                    "running": True,
                    "logged_in": bool(ok),
                    "message": message,
                    "url": url,
                }
            except Exception as e:
                return {
                    "running": True,
                    "logged_in": False,
                    "message": f"照管 Session 檢查失敗：{type(e).__name__}: {e}",
                    "url": url,
                }

        lower = url.lower()
        logged = (
            cfg["host"] in lower
            and "/login" not in lower
            and "/signin" not in lower
            and lower.rstrip("/") != "https://luna.compal-health.com"
        )
        return {
            "running": True,
            "logged_in": logged,
            "message": "已偵測到登入後頁面。" if logged else "等待使用者在官方頁面完成登入。",
            "url": url,
        }

    except Exception as e:
        return {
            "running": cdp_available(system),
            "logged_in": False,
            "message": f"狀態檢查失敗：{type(e).__name__}: {e}",
            "url": "",
        }


LCMS_CASE_URL = (
    "https://csms.mohw.gov.tw/lcms/ca/filter/"
    "?caseno=&serialno=&name=&idno=&pi400aName="
    "&_qd111Resp=&_qd1115Resp=&_qd111adjHint=&_qd111pi400Hint="
    "&_checkAa10Rsp=&_qp110adjHint=&_flexErrHint=&_ca110Modify=&_caseDeath="
    "&_uploadCc01=&_qp300HintPi400=&_qd300HintPi400=&_qd310HintPi400="
    "&_cmsLowerB=&_fh410NotExistsB="
    "&birthDt1=&birthDt2=&applyDt1=&applyDt2=&openDt1=&openDt2="
    "&closeDt1=&closeDt2=&qevalDt1=&qevalDt2=&qmaxInstructDt1=&qmaxInstructDt2="
    "&processDt1=&processD2=&applySource=&_hpCreated=&sextype=&applyType=&censusDiff="
    "&aborigine=&raceType=&liveType=&isdisbook=&sptype=&cmsLev=&disLev=&levcode="
    "&isSick=&discode=&qcntcode=&twnspcode=&vilgcode=&qinformCntcode="
    "&informTwnspcode=&informVilgcode=&qQd120ServDt_b=&qQd120ServDt_e="
    "&sb210id=&sb500id=&ca113exists=&qca113title=&qservUser=&sb400id="
    "&doQuery=true&qdList1=yes&limit=100&offset=0&order=asc"
)

LCMS_QD120_URL = (
    "https://csms.mohw.gov.tw/lcms/qd/filterQd120A/{case_id}"
    "?doQuery=yes&ca100id={case_id}&perms=true&stype=&sourceType=&status="
    "&servDt1=&servDt2=&qd120APi400=&servUserName=&aa10Status="
    "&limit=100&offset=0&order=asc"
)

def _system_page(system):
    cfg = SYSTEMS[system]
    targets = cdp_json(system, "/json/list")
    pages = [
        x for x in targets
        if isinstance(x, dict)
        and x.get("type") == "page"
        and cfg["host"] in str(x.get("url") or "")
    ]
    return pages[0] if pages else None

def _cdp_call(ws_url, method, params=None, call_id=1, timeout=30):
    ws = websocket.create_connection(ws_url, timeout=timeout, suppress_origin=True)
    try:
        ws.send(json.dumps({
            "id": call_id,
            "method": method,
            "params": params or {},
        }))
        while True:
            payload = json.loads(ws.recv())
            if payload.get("id") == call_id:
                if "error" in payload:
                    raise RuntimeError(str(payload["error"]))
                return payload.get("result", {})
    finally:
        try:
            ws.close()
        except Exception:
            pass

def build_cdp_session(system):
    page = _system_page(system)
    if not page or not page.get("webSocketDebuggerUrl"):
        raise RuntimeError("找不到已開啟的官方頁面。")

    ws_url = page["webSocketDebuggerUrl"]
    cookie_result = _cdp_call(ws_url, "Storage.getCookies", call_id=101, timeout=15)
    ua_result = _cdp_call(
        ws_url,
        "Runtime.evaluate",
        {"expression": "navigator.userAgent", "returnByValue": True},
        call_id=102,
        timeout=15,
    )

    s = requests.Session()
    for ck in cookie_result.get("cookies", []) or []:
        try:
            s.cookies.set(
                ck.get("name", ""),
                ck.get("value", ""),
                domain=ck.get("domain") or None,
                path=ck.get("path") or "/",
            )
        except Exception:
            pass

    ua = ua_result.get("result", {}).get("value") or "Mozilla/5.0"
    s.headers.update({
        "User-Agent": ua,
        "Accept": "application/json, text/javascript, */*; q=0.01",
        "X-Requested-With": "XMLHttpRequest",
        "Referer": SYSTEMS[system]["url"],
    })
    return s

def browser_fetch_json(system, url, timeout=30):
    page = _system_page(system)
    if not page or not page.get("webSocketDebuggerUrl"):
        raise RuntimeError(f"{SYSTEMS[system]['name']} 官方頁面尚未開啟。")

    expression = (
        "(async () => {"
        "  const r = await fetch(" + json.dumps(url) + ", {"
        "    method: 'GET',"
        "    credentials: 'include',"
        "    cache: 'no-store',"
        "    headers: {"
        "      'Accept': 'application/json, text/javascript, */*; q=0.01',"
        "      'X-Requested-With': 'XMLHttpRequest'"
        "    }"
        "  });"
        "  const text = await r.text();"
        "  return {status:r.status, ok:r.ok, contentType:r.headers.get('content-type') || '', text:text};"
        "})()"
    )

    result = _cdp_call(
        page["webSocketDebuggerUrl"],
        "Runtime.evaluate",
        {
            "expression": expression,
            "awaitPromise": True,
            "returnByValue": True,
        },
        call_id=205,
        timeout=max(15, timeout),
    )

    payload = result.get("result", {}).get("value")
    if not isinstance(payload, dict):
        raise RuntimeError("瀏覽器未回傳有效資料。")

    status = int(payload.get("status") or 0)
    if status in (401, 403):
        raise RuntimeError(f"照管登入資料已失效 (HTTP {status})")
    if not payload.get("ok"):
        raise RuntimeError(f"照管查詢失敗 (HTTP {status})")

    text = payload.get("text") or ""
    try:
        return json.loads(text)
    except Exception:
        raise RuntimeError(
            f"照管回傳非 JSON (Content-Type={payload.get('contentType','')})"
        )


def safe_json_get(session, url, tag, timeout=30):
    try:
        response = session.get(url, timeout=timeout)
        if response.status_code in (401, 403):
            raise RuntimeError(
                f"{tag}：照管登入資料已失效 (HTTP {response.status_code})"
            )
        response.raise_for_status()
        try:
            return response.json()
        except Exception:
            raise RuntimeError(
                f"{tag}：照管回傳非 JSON，可能已回到登入頁。"
            )

    except requests.exceptions.SSLError:
        # Python/OpenSSL 對照管憑證鏈驗證失敗時，
        # 改由已登入的 Chrome 執行同源 fetch。
        # 這不是關閉 SSL 驗證，仍由 Chrome 正常驗證 HTTPS。
        return browser_fetch_json("lcms", url, timeout=timeout)

def probe_lcms():
    if not cdp_available("lcms"):
        return False, "照管專用 Chrome 尚未啟動。"
    try:
        session = build_cdp_session("lcms")
        probe_url = re.sub(r"limit=\d+", "limit=1", LCMS_CASE_URL)
        data = safe_json_get(session, probe_url, "照管狀態", timeout=10)
        if isinstance(data, dict) and ("rows" in data or "total" in data):
            return True, "照管 Session 有效。"
        return False, "照管頁面已開啟，但尚未取得有效 Session。"
    except Exception as e:
        return False, str(e)

def _safe_total(value, fallback):
    try:
        text = str(value).replace(",", "").strip()
        return int(text)
    except Exception:
        return fallback

def fetch_all_lcms_cases(session):
    all_rows = []
    offset = 0

    while True:
        url = re.sub(r"offset=\d+", f"offset={offset}", LCMS_CASE_URL)
        data = browser_fetch_json("lcms", url, timeout=90)
        if not isinstance(data, dict):
            raise RuntimeError("CA_FILTER 回傳格式不是物件。")

        batch = data.get("rows", []) or []
        if not isinstance(batch, list):
            raise RuntimeError("CA_FILTER rows 格式異常。")

        total = _safe_total(data.get("total"), len(batch))
        all_rows.extend(batch)

        if len(all_rows) >= total or len(batch) < 100:
            break

        offset += 100
        if offset > 5000:
            break

    result = []
    seen = set()
    for row in all_rows:
        if not isinstance(row, dict):
            continue
        case_id = row.get("id")
        if not case_id:
            continue
        try:
            case_id = int(case_id)
        except Exception:
            continue
        if case_id in seen:
            continue
        seen.add(case_id)
        result.append({
            "id": case_id,
            "name": str(row.get("name") or "(未取得姓名)").strip(),
        })
    return result

def to_roc_encoded_date(d):
    return f"{d.year - 1911}%2F{d.month:02d}%2F{d.day:02d}"

def query_qd120_rows(session, case_id, start_date, end_date=None):
    base_url = LCMS_QD120_URL.format(case_id=case_id)
    end_date = end_date or start_date
    roc_start = to_roc_encoded_date(start_date)
    roc_end = to_roc_encoded_date(end_date)
    url0 = (
        base_url
        .replace("servDt1=&", f"servDt1={roc_start}&")
        .replace("servDt2=&", f"servDt2={roc_end}&")
    )

    all_rows = []
    offset = 0

    while True:
        url = re.sub(r"offset=\d+", f"offset={offset}", url0)
        data = browser_fetch_json("lcms", url, timeout=60)
        if not isinstance(data, dict):
            raise RuntimeError("QD120A 回傳格式不是物件。")

        batch = data.get("rows", []) or []
        if not isinstance(batch, list):
            raise RuntimeError("QD120A rows 格式異常。")

        total = _safe_total(data.get("total"), len(batch))
        all_rows.extend(batch)

        if len(all_rows) >= total or len(batch) < 100:
            break

        offset += 100
        if offset > 5000:
            break

    return all_rows

def clean_text(value):
    return re.sub(r"<.*?>", "", str(value or "")).strip()

def parse_service_datetime(serv_dt, hhmm):
    try:
        date_text = str(serv_dt or "").strip()
        if "/" in date_text:
            y, m, d = date_text.split("/")[:3]
        elif "-" in date_text:
            y, m, d = date_text.split("-")[:3]
        else:
            return None, None

        y = int(y)
        if y < 200:
            y += 1911

        raw = re.sub(r"\[.*?\]", "", clean_text(hhmm))
        raw = raw.replace("～", "~").replace("–", "-").replace("—", "-")

        if "~" in raw:
            t1, t2 = [x.strip() for x in raw.split("~", 1)]
        elif "-" in raw:
            t1, t2 = [x.strip() for x in raw.split("-", 1)]
        else:
            return None, None

        start = datetime.strptime(
            f"{y:04d}/{int(m):02d}/{int(d):02d} {t1}",
            "%Y/%m/%d %H:%M",
        )
        end = datetime.strptime(
            f"{y:04d}/{int(m):02d}/{int(d):02d} {t2}",
            "%Y/%m/%d %H:%M",
        )

        if end <= start:
            end += timedelta(days=1)

        return start, end
    except Exception:
        return None, None

def _first_value(row, keys, default=""):
    for key in keys:
        value = row.get(key)
        if value is not None and str(value).strip():
            return str(value).strip()
    return default

def normalize_qd_row(case_id, case_name, row):
    if not isinstance(row, dict):
        return None

    start, end = parse_service_datetime(
        row.get("servDt", ""),
        row.get("hhmm", ""),
    )
    if not start or not end:
        return None

    title = str(row.get("title") or "").strip()
    code_match = re.match(r"([A-Z0-9]{3,5})", title)
    code = code_match.group(1) if code_match else title

    return {
        "case_id": case_id,
        "client": case_name,
        "unit": _first_value(
            row,
            [
                "orgName", "instName", "companyName", "sb400Name",
                "sb400idName", "agencyName", "unitName", "providerName",
            ],
            "(未取得機構)",
        ),
        "worker": _first_value(
            row,
            [
                "servUserName", "servUser", "servUserNm",
                "servUserFullName", "userName", "workerName",
                "caregiverName", "empName",
            ],
            "(未取得居服員)",
        ),
        "worker_id": _first_value(
            row,
            [
                "servUserId", "servUserID", "servUserNo",
                "employeeNo", "workerId", "empId", "sb210id",
                "userId", "staffId",
            ],
            "",
        ),
        "title": title,
        "code": code,
        "date": start.strftime("%Y-%m-%d"),
        "start": start.strftime("%H:%M"),
        "end": end.strftime("%H:%M"),
        "_start_dt": start,
        "_end_dt": end,
    }

def analyze_services(rows):
    issues = []
    by_case = {}

    for row in rows:
        key = (row["case_id"], row["date"])
        by_case.setdefault(key, []).append(row)

    for (_, service_date), items in by_case.items():
        items.sort(key=lambda x: x["_start_dt"])

        for i, a in enumerate(items):
            for b in items[i + 1:]:
                if b["_start_dt"] >= a["_end_dt"]:
                    break
                if not (
                    a["_start_dt"] < b["_end_dt"]
                    and b["_start_dt"] < a["_end_dt"]
                ):
                    continue

                overlap_start = max(a["_start_dt"], b["_start_dt"])
                overlap_end = min(a["_end_dt"], b["_end_dt"])
                minutes = max(
                    1,
                    int((overlap_end - overlap_start).total_seconds() // 60),
                )

                issues.append({
                    "type": "case_overlap",
                    "title": "服務時間重疊",
                    "client": a["client"],
                    "date": service_date,
                    "minutes": minutes,
                    "a": {
                        k: a[k]
                        for k in ("unit", "worker", "title", "start", "end")
                    },
                    "b": {
                        k: b[k]
                        for k in ("unit", "worker", "title", "start", "end")
                    },
                })

    return issues, {
        "total_services": len(rows),
        "total_issues": len(issues),
    }


def normalize_case_name(value):
    text = str(value or "").strip()
    text = re.sub(r"\s+", "", text)
    return text

def parse_case_list_file(file_storage):
    filename = str(file_storage.filename or "").lower()
    raw = file_storage.read()

    names = []

    if filename.endswith((".xlsx", ".xlsm")):
        from io import BytesIO
        wb = load_workbook(BytesIO(raw), read_only=True, data_only=True)
        try:
            ws = wb.worksheets[0]
            for row in ws.iter_rows(min_col=1, max_col=1, values_only=True):
                name = normalize_case_name(row[0] if row else "")
                if name and name not in ("姓名", "個案姓名", "個案"):
                    names.append(name)
        finally:
            wb.close()

    elif filename.endswith((".csv", ".txt")):
        text = None
        for enc in ("utf-8-sig", "utf-8", "cp950", "big5"):
            try:
                text = raw.decode(enc)
                break
            except Exception:
                pass
        if text is None:
            raise ValueError("無法辨識文字檔編碼。")

        for line in text.splitlines():
            first = line.split(",", 1)[0].split("\t", 1)[0]
            name = normalize_case_name(first)
            if name and name not in ("姓名", "個案姓名", "個案"):
                names.append(name)
    else:
        raise ValueError("僅支援 .xlsx、.xlsm、.csv、.txt。")

    result = []
    seen = set()
    for name in names:
        if name in seen:
            continue
        seen.add(name)
        result.append(name)

    return result

@app.post("/case-list/parse")
@require_auth
def parse_case_list():
    file_storage = request.files.get("file")
    if not file_storage or not file_storage.filename:
        return jsonify(error="請選擇個案名單檔案。"), 400

    try:
        names = parse_case_list_file(file_storage)
    except Exception as e:
        return jsonify(error=f"名單讀取失敗：{type(e).__name__}: {e}"), 400

    if not names:
        return jsonify(error="檔案中沒有讀取到個案姓名。"), 400

    return jsonify(ok=True, names=names, count=len(names))

@app.get("/health")
def health():
    users = load_users()
    return jsonify(
        ok=True,
        version=VERSION,
        auth=True,
        admin=True,
        configured=bool(users),
    )

@app.post("/connect/<system>")
@require_auth
def connect(system):
    if system not in SYSTEMS:
        return jsonify(error="unknown system"), 404
    try:
        started = launch_login(system)
    except Exception as e:
        return jsonify(error=str(e)), 500

    return jsonify(
        ok=True,
        started=started,
        message="已開啟官方登入頁，請完成登入。" if started else "專用 Chrome 已經開啟，請在該視窗完成登入。",
    )

@app.get("/status/<system>")
@require_auth
def status(system):
    if system not in SYSTEMS:
        return jsonify(error="unknown system"), 404

    try:
        s = browser_status(system)
        return jsonify(
            system=system,
            logged_in=bool(s.get("logged_in")),
            running=bool(s.get("running")),
            message=str(s.get("message") or ""),
            url=str(s.get("url") or ""),
        )
    except Exception as e:
        # 狀態檢查永遠回傳可讀 JSON，不讓前端只看到 Flask 500 HTML。
        return jsonify(
            system=system,
            logged_in=False,
            running=cdp_available(system),
            message=f"狀態檢查失敗：{type(e).__name__}: {e}",
            url="",
        ), 200

def demo_rows():
    return [
        {"unit":"大安心","worker":"王小明","client":"陳OO","date":"2026-10-06","start":"09:00","end":"10:00"},
        {"unit":"大慶","worker":"王小明","client":"李OO","date":"2026-10-06","start":"09:30","end":"10:30"},
        {"unit":"大安心","worker":"李小天","client":"林OO","date":"2026-10-06","start":"11:00","end":"12:00"},
    ]

@app.get("/demo")
@require_auth
def demo():
    return jsonify(rows=demo_rows())

@app.post("/services")
@require_auth
def services():
    payload = request.get_json(silent=True) or {}
    mode = str(payload.get("mode") or "day").strip().lower()

    try:
        if mode == "month":
            raw_month = str(payload.get("month") or "").strip()
            month_start = datetime.strptime(raw_month, "%Y-%m").date().replace(day=1)
            if month_start.month == 12:
                next_month = month_start.replace(year=month_start.year + 1, month=1)
            else:
                next_month = month_start.replace(month=month_start.month + 1)
            selected_start = month_start
            selected_end = next_month - timedelta(days=1)
        else:
            raw_date = payload.get("date") or str(date.today())
            selected_start = datetime.strptime(raw_date, "%Y-%m-%d").date()
            selected_end = selected_start
            mode = "day"
    except Exception:
        return jsonify(
            error="日期或月份格式錯誤",
            stage="date",
        ), 400

    try:
        ok, message = probe_lcms()
        if not ok:
            return jsonify(
                error=f"照管尚未登入或 Session 已失效：{message}",
                stage="login",
            ), 409

        session = build_cdp_session("lcms")

        try:
            cases = fetch_all_lcms_cases(session)

            requested_case_names = payload.get("case_names") or []
            requested_case_names = [
                normalize_case_name(x)
                for x in requested_case_names
                if normalize_case_name(x)
            ]

            unmatched_case_names = []
            if requested_case_names:
                by_name = {}
                for info in cases:
                    key = normalize_case_name(info.get("name"))
                    by_name.setdefault(key, []).append(info)

                selected_cases = []
                seen_case_ids = set()
                for requested_name in requested_case_names:
                    matches = by_name.get(requested_name, [])
                    if not matches:
                        unmatched_case_names.append(requested_name)
                        continue
                    for info in matches:
                        if info["id"] in seen_case_ids:
                            continue
                        seen_case_ids.add(info["id"])
                        selected_cases.append(info)

                cases = selected_cases
            else:
                unmatched_case_names = []
        except Exception as e:
            return jsonify(
                error=f"取得照管個案清單失敗：{type(e).__name__}: {e}",
                stage="case_list",
            ), 502

        normalized = []
        failures = []

        for info in cases:
            try:
                qd_rows = query_qd120_rows(
                    session,
                    info["id"],
                    selected_start,
                    selected_end,
                )

                for row in qd_rows:
                    item = normalize_qd_row(
                        info["id"],
                        info["name"],
                        row,
                    )
                    if item:
                        normalized.append(item)

            except Exception as e:
                failures.append({
                    "case": info.get("name", ""),
                    "case_id": info.get("id", ""),
                    "error": f"{type(e).__name__}: {e}",
                })

        try:
            issues, stats = analyze_services(normalized)
        except Exception as e:
            return jsonify(
                error=f"衝突分析失敗：{type(e).__name__}: {e}",
                stage="analysis",
                rows_read=len(normalized),
            ), 500

        stats["cases_scanned"] = len(cases)
        stats["case_query_failures"] = len(failures)
        stats["requested_cases"] = len(requested_case_names)
        stats["unmatched_cases"] = len(unmatched_case_names)
        stats["range_mode"] = mode
        stats["range_start"] = selected_start.isoformat()
        stats["range_end"] = selected_end.isoformat()

        public_rows = [
            {
                k: v
                for k, v in row.items()
                if not k.startswith("_")
            }
            for row in normalized
        ]

        return jsonify(
            rows=public_rows,
            issues=issues,
            stats=stats,
            query_failures=failures[:20],
            unmatched_case_names=unmatched_case_names[:100],
            source="LCMS QD120A via Chrome",
        )

    except Exception as e:
        return jsonify(
            error=f"Connector 執行失敗：{type(e).__name__}: {e}",
            stage="unexpected",
        ), 500

def run_server():
    app.run(
        host="127.0.0.1",
        port=PORT,
        debug=False,
        use_reloader=False,
        threaded=True,
    )

def show_window():
    root = tk.Tk()
    root.title("服務衝突連線器")
    root.geometry("420x220")
    root.resizable(False, False)

    tk.Label(root, text="服務衝突連線器", font=("Microsoft JhengHei UI", 18, "bold")).pack(pady=(22, 8))
    tk.Label(
        root,
        text=f"Connector 已啟動 v{VERSION}\nhttp://127.0.0.1:{PORT}",
        font=("Microsoft JhengHei UI", 11),
        justify="center",
    ).pack(pady=8)

    btn_frame = tk.Frame(root)
    btn_frame.pack(pady=14)

    tk.Button(
        btn_frame,
        text="開啟服務衝突網站",
        width=18,
        command=lambda: webbrowser.open(WEBSITE_URL),
    ).pack(side="left", padx=6)

    tk.Button(
        btn_frame,
        text="關閉連線器",
        width=14,
        command=root.destroy,
    ).pack(side="left", padx=6)

    tk.Label(
        root,
        text="請保持此連線器開啟，再回到網頁登入。",
        fg="#666666",
        font=("Microsoft JhengHei UI", 9),
    ).pack(pady=(8, 0))

    root.mainloop()

def running_connector_version():
    try:
        data = requests.get(
            f"http://127.0.0.1:{PORT}/health",
            timeout=1.5,
        ).json()
        return str(data.get("version") or "未知")
    except Exception:
        return "未知"


def main():
    if already_running():
        running_version = running_connector_version()

        if running_version != VERSION:
            messagebox.showwarning(
                "偵測到舊版 Connector",
                (
                    f"目前電腦上已有 Connector v{running_version} 正在執行。\n\n"
                    f"你剛開啟的是 v{VERSION}，但 127.0.0.1:{PORT} "
                    "仍被舊版占用，因此網頁實際連到的仍是舊版。\n\n"
                    "請先關閉舊的「服務衝突連線器」視窗，"
                    "再重新開啟最新版。"
                ),
            )
        else:
            messagebox.showinfo(
                "服務衝突連線器",
                f"Connector v{VERSION} 已經在執行中。",
            )
        return

    t = threading.Thread(target=run_server, daemon=True)
    t.start()
    show_window()

if __name__ == "__main__":
    main()
